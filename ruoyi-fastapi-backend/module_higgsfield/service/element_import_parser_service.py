"""元素表格解析为待复核候选，保留定位与错误，不创建正式元素。"""
import csv
import io
from dataclasses import dataclass
from types import SimpleNamespace
from openpyxl import load_workbook
from pydantic import ValidationError
from exceptions.exception import ServiceWarning
from module_higgsfield.entity.vo.element_vo import ElementFields
from module_higgsfield.service.document_storage_service import DocumentStorageService


@dataclass(frozen=True)
class ElementImportCandidate:
    row_no: int
    source_locator: dict
    candidate: dict
    issues: list[dict]


class ElementImportParserService:
    version = 'element-table-v1'
    columns = ('name', 'alias', 'category', 'description')
    chinese_columns = ('名称', '别名', '分类', '描述')
    xlsx_mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    categories = {'角色': 'character', '场景': 'scene', '道具': 'prop', '其他': 'other'}

    @staticmethod
    def fail(line, message):
        raise ServiceWarning(message=message, data={'errorCode': 'ELEMENT_IMPORT_FILE_INVALID', 'lineNo': line})

    @classmethod
    def limits(cls, body, max_bytes, max_rows):
        if (type(max_bytes) is not int or max_bytes <= 0 or type(max_rows) is not int
                or not 1 <= max_rows <= 10000 or not isinstance(body, bytes) or not body or len(body) > max_bytes):
            cls.fail(0, '导入文件为空或大小、候选数量限制无效')

    @classmethod
    def header(cls, values):
        if tuple(values) not in (cls.columns, cls.chinese_columns):
            cls.fail(1, '表头必须依次为名称、别名、分类、描述，或对应英文模板字段')

    @classmethod
    def candidate(cls, values, row_no, locator, *, project_name, include_project_name):
        if len(values) != 4 or any(not isinstance(value, str) for value in values):
            cls.fail(locator.get('line', locator.get('row', 0)), '模板行必须包含四个文本字段，不能使用数值或公式')
        if any(len(value) > 20000 or any(ord(char) < 32 and char not in '\t\r\n' for char in value) for value in values):
            cls.fail(locator.get('line', locator.get('row', 0)), '候选字段过长或包含不支持的控制字符')
        name, alias, category, description = values
        alias = alias.strip()
        if include_project_name and alias:
            alias = project_name.strip() + '_' + alias
        candidate = dict(schemaVersion=1, name=name, alias=alias,
                         category=cls.categories.get(category.strip(), category.strip()),
                         description=description, propertyValues=[])
        issues = []
        try:
            checked = ElementFields.model_validate({key: value for key, value in candidate.items() if key != 'schemaVersion'})
            candidate.update(name=checked.name, alias=checked.alias)
        except ValidationError as error:
            # 不把Pydantic输入正文或异常上下文带入诊断。
            issues = [{'field': str(item['loc'][0]), 'code': 'ELEMENT_FIELD_INVALID',
                       'message': '字段不符合元素规范，请修正后确认'} for item in error.errors(include_input=False, include_context=False)]
        return ElementImportCandidate(row_no, locator, candidate, issues)

    @classmethod
    def finish(cls, rows):
        if not rows:
            cls.fail(1, '文件没有可复核的元素候选')
        aliases = {}
        for row in rows:
            alias = row.candidate['alias']
            if alias:
                aliases.setdefault(alias, []).append(row)
        for duplicates in aliases.values():
            if len(duplicates) > 1:
                for row in duplicates:
                    row.issues.append({'field': 'alias', 'code': 'ELEMENT_ALIAS_DUPLICATE',
                                       'message': '本批次存在相同别名，请修正或取消选择重复项'})
        return rows

    @staticmethod
    def options(project_name, include_project_name):
        if type(include_project_name) is not bool or not isinstance(project_name, str):
            raise ValueError('项目名称与别名选项类型无效')
        if include_project_name and (not project_name.strip() or any(ord(char) < 32 for char in project_name)):
            raise ValueError('包含项目名时必须提供有效名称')

    @classmethod
    def csv(cls, body, *, max_bytes, max_rows=10000, project_name='', include_project_name=False):
        cls.limits(body, max_bytes, max_rows)
        cls.options(project_name, include_project_name)
        rows = []
        try:
            reader = csv.reader(io.StringIO(body.decode('utf-8-sig'), newline=''), strict=True)
            cls.header(next(reader, []))
            while True:
                line = reader.line_num + 1
                values = next(reader, None)
                if values is None:
                    break
                if not values or all(value == '' for value in values):
                    continue
                if len(rows) >= max_rows:
                    cls.fail(line, '元素候选数量超过上限')
                rows.append(cls.candidate(values, len(rows) + 1, {'line': line},
                    project_name=project_name, include_project_name=include_project_name))
        except (UnicodeError, csv.Error):
            cls.fail(0, 'CSV必须为UTF-8编码且引号及列结构完整')
        return cls.finish(rows)

    @classmethod
    def xlsx(cls, body, *, max_bytes, max_rows=10000, project_name='', include_project_name=False):
        cls.limits(body, max_bytes, max_rows)
        cls.options(project_name, include_project_name)
        workbook = None
        try:
            DocumentStorageService(SimpleNamespace(max_bytes=max_bytes)).verify_file(io.BytesIO(body), cls.xlsx_mime)
            workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
            if len(workbook.worksheets) != 1:
                cls.fail(0, '元素模板必须只有一个工作表')
            sheet = workbook.worksheets[0]
            if sheet.sheet_state != 'visible' or sheet.max_column != 4:
                cls.fail(1, '元素模板必须为可见的四列表格')
            rows = []
            for number, cells in enumerate(sheet.iter_rows(), 1):
                if any(cell.data_type == 'f' for cell in cells):
                    cls.fail(number, '元素模板不能含公式，请使用文本原始值')
                values = ['' if cell.value is None else cell.value for cell in cells]
                if number == 1:
                    cls.header(values)
                    continue
                if all(value == '' for value in values):
                    continue
                if len(rows) >= max_rows:
                    cls.fail(number, '元素候选数量超过上限')
                rows.append(cls.candidate(values, len(rows) + 1, {'sheet': sheet.title, 'row': number},
                    project_name=project_name, include_project_name=include_project_name))
            return cls.finish(rows)
        except ServiceWarning as error:
            if error.data and error.data.get('errorCode') == 'ELEMENT_IMPORT_FILE_INVALID':
                raise
            cls.fail(0, '工作簿结构无效或超过安全限制')
        except Exception:
            cls.fail(0, '工作簿内容无法读取')
        finally:
            if workbook is not None:
                workbook.close()
