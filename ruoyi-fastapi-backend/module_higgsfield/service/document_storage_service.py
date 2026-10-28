"""管理证据内容校验；只读取数据，不执行公式、宏或外部链接。"""
import csv
import io
import zipfile

from openpyxl import load_workbook
from exceptions.exception import ServiceWarning
from module_higgsfield.service.storage_service import StorageService
from module_higgsfield.service.docx_content_service import DocxContentService
from module_higgsfield.service.pdf_content_service import PdfContentService
from module_higgsfield.service.pptx_content_service import PptxContentService


class DocumentStorageService(StorageService):
    def verify_file(self, handle, declared_mime):
        try:
            handle.seek(0)
            if declared_mime in ('image/png', 'image/jpeg', 'image/webp'):
                # 复用实际解码、像素限制和单帧约束；完整组装仍核对声明 MIME。
                return self.verify_media(handle)
            if declared_mime == PdfContentService.mime:
                PdfContentService.pages(handle.read(self.policy.max_bytes + 1), max_bytes=self.policy.max_bytes)
            elif declared_mime == DocxContentService.mime:
                DocxContentService.paragraphs(handle.read(self.policy.max_bytes + 1), max_bytes=self.policy.max_bytes)
            elif declared_mime == PptxContentService.mime:
                PptxContentService.slides(handle.read(self.policy.max_bytes + 1), max_bytes=self.policy.max_bytes)
            elif declared_mime in ('text/plain', 'text/markdown', 'text/csv', 'text/html'):
                data = handle.read(self.policy.max_bytes + 1)
                if len(data) > self.policy.max_bytes:
                    raise ValueError
                content = data.decode('utf-8-sig')
                if not content.strip() or any(ord(c) < 32 and c not in '\t\r\n' for c in content):
                    raise ValueError
                # HTML只校验文本编码；后续必须以文本或sandbox预览，不能直接执行。
                if declared_mime == 'text/csv':
                    width, count = None, 0
                    for row in csv.reader(io.StringIO(content, newline=''), strict=True):
                        if not row:
                            continue
                        width = len(row) if width is None else width
                        count += 1
                        if len(row) != width or width > 1000 or count > 100000:
                            raise ValueError
                    if not count:
                        raise ValueError
            elif declared_mime == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
                # 压缩包先限展开体积和成员数，禁止宏及外链，再实际读取工作表。
                with zipfile.ZipFile(handle) as archive:
                    members = archive.infolist()
                    names = [item.filename for item in members]
                    if (len(members) > 10000 or len(names) != len(set(names))
                            or sum(item.file_size for item in members) > self.policy.max_bytes * 10
                            or any('vbaproject' in name.lower() or 'externallinks/' in name.lower() for name in names)):
                        raise ValueError
                    for item in members:
                        if item.filename.lower().endswith(('.xml', '.rels')):
                            content = archive.read(item).replace(b'\x00', b'').upper()
                            if b'<!DOCTYPE' in content or b'<!ENTITY' in content:
                                raise ValueError
                handle.seek(0)
                workbook = load_workbook(handle, read_only=True, data_only=True, keep_links=False)
                try:
                    cells = 0
                    if not workbook.worksheets or len(workbook.worksheets) > 100:
                        raise ValueError
                    for sheet in workbook.worksheets:
                        if (sheet.max_column or 0) > 1000 or (sheet.max_row or 0) > 100000:
                            raise ValueError
                        for row in sheet.iter_rows():
                            cells += len(row)
                            if cells > 1000000:
                                raise ValueError
                finally:
                    workbook.close()
            else:
                raise ValueError
        except Exception as error:
            if isinstance(error, ServiceWarning):
                raise
            raise ServiceWarning(message='文档格式损坏、编码不支持或超出结构限制',
                data={'errorCode': 'MEDIA_INVALID'}) from None
        return {'mime_type': declared_mime, 'width': None, 'height': None, 'seconds': None}
