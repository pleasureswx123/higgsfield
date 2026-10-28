"""标准账单模板解析；不猜测供应商字段、不推断币种或计费对象。"""
import csv
import io
import re
from dataclasses import dataclass
from decimal import Decimal
from types import SimpleNamespace
from openpyxl import load_workbook
from module_higgsfield.service.document_storage_service import DocumentStorageService
from exceptions.exception import ServiceWarning


@dataclass(frozen=True)
class ReconciliationRow:
    line_no: int
    external_line_key: str
    currency: str
    amount: Decimal
    raw_line: dict


class ReconciliationParserService:
    columns = ('externalLineKey', 'currency', 'amount')
    xlsx_mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    mime_types = ('text/csv', xlsx_mime)

    @classmethod
    def xlsx(cls, body, max_bytes, max_rows=100000):
        """标准单表模板仅接受文本单元格，避免 Excel 数值精度损失与公式缓存歧义。"""
        if not isinstance(body, bytes) or not body or len(body) > max_bytes or max_rows <= 0:
            cls.fail(0, '账单文件为空或超过导入大小限制')
        try:
            DocumentStorageService(SimpleNamespace(max_bytes=max_bytes)).verify_file(io.BytesIO(body), cls.xlsx_mime)
        except ServiceWarning:
            cls.fail(0, '账单工作簿结构无效或超过限制')
        workbook = None
        try:
            workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=False, keep_links=False)
            if len(workbook.worksheets) != 1:
                cls.fail(0, '账单 XLSX 必须仅有一个工作表')
            sheet = workbook.worksheets[0]
            if sheet.sheet_state != 'visible' or sheet.max_column != 3:
                cls.fail(1, '账单工作表必须可见且仅有三个模板列')
            rows = []
            for line_no, cells in enumerate(sheet.iter_rows(), 1):
                values = [cell.value for cell in cells]
                if any(cell.data_type == 'f' for cell in cells):
                    cls.fail(line_no, '账单不接受公式，请使用文本原始值')
                if line_no == 1:
                    if tuple(values) != cls.columns:
                        cls.fail(1, '账单表头必须依次为 externalLineKey,currency,amount')
                    continue
                if all(value is None for value in values):
                    continue
                if len(rows) >= max_rows:
                    cls.fail(line_no, '账单行数超过导入限制')
                rows.append(cls.row(line_no, values))
            if not rows:
                cls.fail(1, '账单没有可导入明细')
            return rows
        except ServiceWarning:
            raise
        except Exception:
            cls.fail(0, '账单工作簿内容无法读取')
        finally:
            if workbook is not None:
                workbook.close()

    @staticmethod
    def fail(line_no, message):
        raise ServiceWarning(message=message, data={'errorCode': 'RECONCILIATION_FILE_INVALID', 'lineNo': line_no})

    @classmethod
    def row(cls, line_no, values):
        if len(values) != 3 or any(not isinstance(value, str) for value in values):
            cls.fail(line_no, '账单行必须包含稳定账单项标识、币种和文本金额')
        key, currency, amount = values
        if not key.strip() or key != key.strip() or len(key) > 255 or any(ord(char) < 32 for char in key):
            cls.fail(line_no, '账单项标识不能为空、包含控制字符或首尾空白')
        if not re.fullmatch('[A-Z]{3}', currency):
            cls.fail(line_no, '账单币种必须为三个大写字母')
        # 严格匹配 numeric(24,10)，拒绝浮点、指数、千分位和静默舍入。
        if not re.fullmatch(r'-?(?:0|[1-9][0-9]{0,13})(?:\.[0-9]{1,10})?', amount):
            cls.fail(line_no, '金额必须为最多14位整数、10位小数的十进制文本，可负数退款')
        return ReconciliationRow(line_no, key, currency, Decimal(amount), dict(zip(cls.columns, values)))

    @classmethod
    def csv(cls, body, max_bytes, max_rows=100000):
        if not isinstance(body, bytes) or not body or len(body) > max_bytes or max_bytes <= 0 or max_rows <= 0:
            cls.fail(0, '账单文件为空或超过导入大小限制')
        try:
            content = body.decode('utf-8-sig')
        except UnicodeError:
            cls.fail(0, '账单 CSV 必须使用 UTF-8 编码')
        reader = csv.reader(io.StringIO(content, newline=''), strict=True)
        rows = []
        try:
            if tuple(next(reader, [])) != cls.columns:
                cls.fail(1, '账单表头必须依次为 externalLineKey,currency,amount')
            while True:
                line_no = reader.line_num + 1
                values = next(reader, None)
                if values is None:
                    break
                if not values:
                    continue
                if len(rows) >= max_rows:
                    cls.fail(line_no, '账单行数超过导入限制')
                rows.append(cls.row(line_no, values))
        except csv.Error:
            cls.fail(reader.line_num, '账单 CSV 引号或字段结构无效')
        if not rows:
            cls.fail(1, '账单没有可导入明细')
        return rows
